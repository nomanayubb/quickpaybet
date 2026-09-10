import io
from decimal import Decimal
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.bets.models import Bet
from apps.bets.services import place_bet
from apps.casino.models import CasinoBrand, CasinoConfig, CasinoGame
from apps.casino.providers import MockCasinoProvider
from apps.cashback.models import CashbackCredit
from apps.exchange.models import ExchangeFill, ExchangeOrder
from apps.payments.models import CryptoPayment
from apps.sports.models import Match, RealtimeOddsConfig, RefreshMode, Sport, UserMatchOddsOverride
from apps.sports.pricing import (
    normalize_odds, resolve_user_extra_adjustment, resolve_user_extra_lay_spread,
)
from apps.wallet.models import Wallet
from apps.wallet.services import deposit_funds

User = get_user_model()


def _read_xlsx_rows(response):
    """Returns (header_row, data_rows) from an .xlsx HttpResponse's content."""
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    return all_rows[0], all_rows[1:]


class ExcelExportTestsBase(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email='exportadmin@example.com', password='testpass123', is_staff=True,
        )


class ManageUsersExportTests(ExcelExportTestsBase):
    def setUp(self):
        super().setUp()
        self.matching = User.objects.create_user(email='alice@example.com', password='testpass123')
        self.other = User.objects.create_user(email='bob@example.com', password='testpass123')

    def test_filtered_export_only_includes_matching_users(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_users'), {'export': 'users_xlsx', 'q': 'alice'})
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment', response['Content-Disposition'])
        header, rows = _read_xlsx_rows(response)
        emails = [row[0] for row in rows]
        self.assertIn('alice@example.com', emails)
        self.assertNotIn('bob@example.com', emails)

    def test_master_export_ignores_current_filters(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_users'), {'export': 'users_master_xlsx', 'q': 'alice'})
        header, rows = _read_xlsx_rows(response)
        emails = [row[0] for row in rows]
        # A search for "alice" is active in the query string but must have no
        # effect on the master export - it should still include bob.
        self.assertIn('alice@example.com', emails)
        self.assertIn('bob@example.com', emails)

    def test_master_export_never_includes_password_hash(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_users'), {'export': 'users_master_xlsx'})
        header, rows = _read_xlsx_rows(response)
        self.assertNotIn('Password', header)
        for row in rows:
            for cell in row:
                if isinstance(cell, str):
                    self.assertNotIn('pbkdf2', cell.lower())


class MasterExportAccessBoundaryTests(TestCase):
    def test_master_role_admin_export_only_includes_own_children(self):
        """
        The single most important guarantee of this feature: a Master-role
        admin's "master" export must only ever contain their own downstream
        users, never the whole platform - reopening that boundary would be
        the same class of privilege-escalation bug already fixed once in
        this project.
        """
        master = User.objects.create_user(
            email='master1@example.com', password='testpass123', role=User.Role.MASTER,
        )
        child = User.objects.create_user(
            email='child1@example.com', password='testpass123', parent=master,
        )
        unrelated = User.objects.create_user(
            email='unrelated@example.com', password='testpass123',
        )

        self.client.force_login(master)
        response = self.client.get(reverse('web:admin_users'), {'export': 'users_master_xlsx'})
        header, rows = _read_xlsx_rows(response)
        emails = [row[0] for row in rows]
        self.assertIn('child1@example.com', emails)
        self.assertNotIn('unrelated@example.com', emails)
        self.assertNotIn('master1@example.com', emails)


class LedgerExportTests(ExcelExportTestsBase):
    def setUp(self):
        super().setUp()
        self.target = User.objects.create_user(email='ledgertarget@example.com', password='testpass123')
        Wallet.objects.get_or_create(user=self.target)
        deposit_funds(self.target, Decimal('1000'))
        sport = Sport.objects.create(name='LedgerExportSport', slug='ledger-export-sport')
        self.match = Match.objects.create(
            sport=sport, home_team='LX Home', away_team='LX Away',
            start_time=timezone.now() + timezone.timedelta(hours=1), status=Match.Status.SCHEDULED,
            odds_home=Decimal('2.00'), odds_draw=Decimal('3.00'), odds_away=Decimal('4.00'),
        )
        self.bet = place_bet(self.target, self.match.id, 'home', Decimal('10'))
        CryptoPayment.objects.create(
            user=self.target, amount=Decimal('50'), currency='USDT',
            payment_type=CryptoPayment.PaymentType.DEPOSIT, status=CryptoPayment.Status.COMPLETED,
            provider='mock', external_id='ledgertest-1',
        )

    def test_bets_xlsx_export(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('web:admin_user_ledger', args=[self.target.id]), {'export': 'bets_xlsx'},
        )
        header, rows = _read_xlsx_rows(response)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][2], 'home')
        self.assertEqual(rows[0][3], '2.00')

    def test_payments_xlsx_export(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('web:admin_user_ledger', args=[self.target.id]), {'export': 'payments_xlsx'},
        )
        header, rows = _read_xlsx_rows(response)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], 'deposit')
        self.assertEqual(rows[0][7], 'ledgertest-1')


class CasinoLedgerHistoryTests(ExcelExportTestsBase):
    """The per-user Ledger page's Casino History section - shows which game/provider a win or loss happened on."""

    def setUp(self):
        super().setUp()
        from apps.casino.services import handle_waija_wallet_event

        self.target = User.objects.create_user(email='casinoledgertarget@example.com', password='testpass123')
        Wallet.objects.get_or_create(user=self.target)
        deposit_funds(self.target, Decimal('100'))
        self.brand = CasinoBrand.objects.create(brand_id=1, name='Pragmatic Play', is_active=True)
        self.game = CasinoGame.objects.create(
            game_id=1, game_uid='pragmaticslots/gatesofolympus', brand=self.brand, name='Gates of Olympus', is_active=True,
        )
        username = f'qpb{self.target.id}'
        handle_waija_wallet_event(
            username=username, action='debit', amount=Decimal('5.00'), call_id='ledger-call-1',
            round_id='round-1', game_uid=self.game.game_uid, is_rollback=False, event_type='spin', raw_payload={},
        )
        handle_waija_wallet_event(
            username=username, action='credit', amount=Decimal('12.50'), call_id='ledger-call-2',
            round_id='round-1', game_uid=self.game.game_uid, is_rollback=False, event_type='spin', raw_payload={},
        )

    def test_ledger_page_shows_game_and_provider_name(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_user_ledger', args=[self.target.id]))
        self.assertContains(response, 'Gates of Olympus')
        self.assertContains(response, 'Pragmatic Play')

    def test_casino_xlsx_export_includes_win_and_loss_rows(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('web:admin_user_ledger', args=[self.target.id]), {'export': 'casino_xlsx'},
        )
        header, rows = _read_xlsx_rows(response)
        self.assertEqual(len(rows), 2)
        by_action = {row[2]: row for row in rows}
        self.assertEqual(by_action['debit'][1], 'Gates of Olympus (Pragmatic Play)')
        self.assertEqual(by_action['debit'][3], '5.00000000')
        self.assertEqual(by_action['credit'][3], '12.50000000')

    def test_search_filters_by_game_name(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('web:admin_user_ledger', args=[self.target.id]), {'casino_q': 'nonexistent-game'},
        )
        self.assertContains(response, 'No casino activity matches this search.')


class ExchangeExportTests(ExcelExportTestsBase):
    def setUp(self):
        super().setUp()
        self.user1 = User.objects.create_user(email='exchangeuser1@example.com', password='testpass123')
        self.user2 = User.objects.create_user(email='exchangeuser2@example.com', password='testpass123')
        sport = Sport.objects.create(name='ExchangeExportSport', slug='exchange-export-sport')
        self.match = Match.objects.create(
            sport=sport, home_team='EX Home', away_team='EX Away',
            start_time=timezone.now() + timezone.timedelta(hours=1), status=Match.Status.SCHEDULED,
        )
        self.back_order = ExchangeOrder.objects.create(
            user=self.user1, match=self.match, selection=Bet.Selection.HOME, side=ExchangeOrder.Side.BACK,
            odds=Decimal('2.00'), stake=Decimal('10'), matched_stake=Decimal('10'),
            status=ExchangeOrder.Status.SETTLED,
        )
        self.lay_order = ExchangeOrder.objects.create(
            user=self.user2, match=self.match, selection=Bet.Selection.HOME, side=ExchangeOrder.Side.LAY,
            odds=Decimal('2.00'), stake=Decimal('10'), matched_stake=Decimal('10'),
            status=ExchangeOrder.Status.SETTLED,
        )
        ExchangeFill.objects.create(
            match=self.match, selection=Bet.Selection.HOME, back_order=self.back_order, lay_order=self.lay_order,
            odds=Decimal('2.00'), stake=Decimal('10'), status=ExchangeFill.Status.BACK_WON,
            commission_amount=Decimal('0.50'),
        )

    def test_orders_xlsx_export(self):
        self.client.force_login(self.admin)
        # orders_status defaults to "open" on this page (matching the on-screen
        # default view); pass it explicitly empty to get "any status", since
        # both test orders are SETTLED.
        response = self.client.get(
            reverse('web:admin_exchange'), {'export': 'orders_xlsx', 'orders_status': ''},
        )
        header, rows = _read_xlsx_rows(response)
        emails = [row[0] for row in rows]
        self.assertIn('exchangeuser1@example.com', emails)
        self.assertIn('exchangeuser2@example.com', emails)

    def test_fills_xlsx_export(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_exchange'), {'export': 'fills_xlsx'})
        header, rows = _read_xlsx_rows(response)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][2], 'exchangeuser1@example.com')
        self.assertEqual(rows[0][3], 'exchangeuser2@example.com')
        self.assertEqual(rows[0][7], '0.50000000')


class CashbackExportTests(ExcelExportTestsBase):
    def setUp(self):
        super().setUp()
        self.target = User.objects.create_user(email='cashbackuser@example.com', password='testpass123')
        CashbackCredit.objects.create(
            user=self.target, source_description='Lost single bet #1', loss_amount=Decimal('20'),
            rate_applied=Decimal('10.00'), cashback_amount=Decimal('2'), multiplier_applied=Decimal('3.00'),
            wagering_required=Decimal('6'), wagering_progress=Decimal('0'), deduct_original_on_unlock=False,
            status=CashbackCredit.Status.LOCKED,
        )

    def test_credits_xlsx_export(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_cashback'), {'export': 'credits_xlsx'})
        header, rows = _read_xlsx_rows(response)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 'cashbackuser@example.com')
        self.assertEqual(rows[0][2], '20.00000000')


class AuditLogExportTests(ExcelExportTestsBase):
    def setUp(self):
        super().setUp()
        self.target = User.objects.create_user(email='audituser@example.com', password='testpass123')
        AuditLog.objects.create(
            user=self.target, action='test_action', target_type='Match', target_id='5',
            metadata={'note': 'test'}, ip_address='127.0.0.1',
        )

    def test_audit_xlsx_export(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('web:admin_audit'), {'export': 'xlsx'})
        header, rows = _read_xlsx_rows(response)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], 'test_action')
        self.assertIn('note', rows[0][4])


class RefreshModeTests(TestCase):
    def setUp(self):
        cache.clear()
        RealtimeOddsConfig.objects.filter(pk=1).delete()
        self.config = RealtimeOddsConfig.objects.create(
            pk=1, is_enabled=True, live_refresh_seconds=20,
            # Deliberately high so the VIEWER_GATED default never
            # accidentally triggers with just this test client's own
            # single heartbeat - isolates "is ALWAYS actually bypassing
            # the gate" from "did the gate happen to be satisfied anyway".
            min_viewers_for_realtime=5,
        )
        self.sport = Sport.objects.create(
            name='RefreshModeSport', slug='refresh-mode-sport', provider_key='refresh_mode_sport',
        )
        self.match = Match.objects.create(
            sport=self.sport, home_team='H', away_team='A',
            start_time=timezone.now() + timezone.timedelta(hours=1), status=Match.Status.SCHEDULED,
            odds_home=Decimal('2.00'), odds_draw=Decimal('3.00'), odds_away=Decimal('4.00'),
        )

    def _mock_provider(self):
        provider = MagicMock()
        provider.fetch_matches.return_value = [{
            'home_team': 'H', 'away_team': 'A',
            'start_time': self.match.start_time.isoformat(),
            'odds_home': '2.10', 'odds_draw': '3.10', 'odds_away': '4.10',
        }]
        return provider

    def _expected_refreshed_home_odds(self):
        # maybe_refresh_sport_odds applies normalize_odds's 5% margin
        # before storing - the stored value is never the raw provider
        # quote verbatim, matching this project's other realtime tests.
        home, _, _ = normalize_odds('2.10', '3.10', '4.10', margin=Decimal('0.05'))
        return home

    def test_viewer_gated_default_does_not_refresh_with_too_few_viewers(self):
        with patch('apps.sports.realtime.get_odds_provider', return_value=self._mock_provider()):
            response = self.client.get(reverse('web:match_odds_poll', args=[self.match.id]))
        data = response.json()
        self.match.refresh_from_db()
        self.assertEqual(self.match.odds_home, Decimal('2.00'))  # unchanged - not enough viewers
        self.assertTrue(data['should_continue_polling'])
        self.assertFalse(data['manual_refresh_available'])

    def test_always_mode_refreshes_regardless_of_viewer_count(self):
        self.match.refresh_mode_override = RefreshMode.ALWAYS
        self.match.save(update_fields=['refresh_mode_override'])
        with patch('apps.sports.realtime.get_odds_provider', return_value=self._mock_provider()):
            self.client.get(reverse('web:match_odds_poll', args=[self.match.id]))
        self.match.refresh_from_db()
        self.assertEqual(self.match.odds_home, self._expected_refreshed_home_odds())  # refreshed even with too few viewers

    def test_manual_only_mode_never_auto_refreshes(self):
        self.match.refresh_mode_override = RefreshMode.MANUAL_ONLY
        self.match.save(update_fields=['refresh_mode_override'])
        with patch('apps.sports.realtime.get_odds_provider', return_value=self._mock_provider()):
            response = self.client.get(reverse('web:match_odds_poll', args=[self.match.id]))
        data = response.json()
        self.match.refresh_from_db()
        self.assertEqual(self.match.odds_home, Decimal('2.00'))  # unchanged
        self.assertFalse(data['should_continue_polling'])  # poll loop stops itself, no JS changes needed
        self.assertTrue(data['manual_refresh_available'])

    def test_manual_only_mode_refreshes_on_explicit_request(self):
        self.match.refresh_mode_override = RefreshMode.MANUAL_ONLY
        self.match.save(update_fields=['refresh_mode_override'])
        with patch('apps.sports.realtime.get_odds_provider', return_value=self._mock_provider()):
            response = self.client.get(
                reverse('web:match_odds_poll', args=[self.match.id]), {'refresh': 'manual'},
            )
        self.match.refresh_from_db()
        expected = self._expected_refreshed_home_odds()
        self.assertEqual(self.match.odds_home, expected)
        self.assertEqual(response.json()['odds_home'], str(expected))

    def test_manual_only_mode_manual_refresh_is_still_throttled(self):
        self.match.refresh_mode_override = RefreshMode.MANUAL_ONLY
        self.match.save(update_fields=['refresh_mode_override'])
        provider = self._mock_provider()
        with patch('apps.sports.realtime.get_odds_provider', return_value=provider):
            self.client.get(reverse('web:match_odds_poll', args=[self.match.id]), {'refresh': 'manual'})
            self.client.get(reverse('web:match_odds_poll', args=[self.match.id]), {'refresh': 'manual'})
        provider.fetch_matches.assert_called_once()  # second click in the same window is throttled

    def test_stopped_mode_refuses_even_a_manual_refresh(self):
        self.match.refresh_mode_override = RefreshMode.STOPPED
        self.match.save(update_fields=['refresh_mode_override'])
        with patch('apps.sports.realtime.get_odds_provider', return_value=self._mock_provider()) as mock_get:
            response = self.client.get(
                reverse('web:match_odds_poll', args=[self.match.id]), {'refresh': 'manual'},
            )
        self.match.refresh_from_db()
        self.assertEqual(self.match.odds_home, Decimal('2.00'))
        mock_get.assert_not_called()
        data = response.json()
        self.assertFalse(data['should_continue_polling'])
        self.assertFalse(data['manual_refresh_available'])


class AdminEditMatchViewTests(TestCase):
    """
    Real gap found live: this panel form had its own manually-coded
    +-10.00 bound on odds_adjustment left stale after the model itself
    was widened to +-99.00, and never exposed lay_spread_override or
    refresh_mode_override at all (both Django-admin-only until now).
    """
    def setUp(self):
        self.admin = User.objects.create_user(email='matcheditadmin@example.com', password='testpass123', is_staff=True)
        self.client.force_login(self.admin)
        self.sport = Sport.objects.create(name='MatchEditSport', slug='match-edit-sport')
        self.match = Match.objects.create(
            sport=self.sport, home_team='Home', away_team='Away',
            start_time=timezone.now() + timezone.timedelta(hours=1), status=Match.Status.SCHEDULED,
            odds_home=Decimal('2.00'), odds_draw=Decimal('3.00'), odds_away=Decimal('4.00'),
        )

    def _post(self, **overrides):
        data = {
            'home_team': self.match.home_team,
            'away_team': self.match.away_team,
            'status': self.match.status,
            'start_time': self.match.start_time.strftime('%Y-%m-%dT%H:%M'),
            'odds_home': '2.00', 'odds_draw': '3.00', 'odds_away': '4.00',
            'odds_adjustment': '', 'lay_spread_override': '', 'refresh_mode_override': '',
        }
        data.update(overrides)
        return self.client.post(reverse('web:admin_edit_match', args=[self.match.id]), data)

    def test_odds_adjustment_beyond_old_10_limit_is_accepted(self):
        response = self._post(odds_adjustment='42.20')
        self.assertRedirects(response, reverse('web:admin_matches'))
        self.match.refresh_from_db()
        self.assertEqual(self.match.odds_adjustment, Decimal('42.20'))

    def test_odds_adjustment_beyond_99_is_rejected(self):
        response = self._post(odds_adjustment='99.01')
        self.assertEqual(response.status_code, 200)
        self.match.refresh_from_db()
        self.assertIsNone(self.match.odds_adjustment)

    def test_lay_spread_override_saves(self):
        response = self._post(lay_spread_override='0.50')
        self.assertRedirects(response, reverse('web:admin_matches'))
        self.match.refresh_from_db()
        self.assertEqual(self.match.lay_spread_override, Decimal('0.50'))

    def test_lay_spread_override_out_of_range_is_rejected(self):
        response = self._post(lay_spread_override='6.00')
        self.assertEqual(response.status_code, 200)
        self.match.refresh_from_db()
        self.assertIsNone(self.match.lay_spread_override)

    def test_refresh_mode_override_saves(self):
        response = self._post(refresh_mode_override=RefreshMode.MANUAL_ONLY)
        self.assertRedirects(response, reverse('web:admin_matches'))
        self.match.refresh_from_db()
        self.assertEqual(self.match.refresh_mode_override, RefreshMode.MANUAL_ONLY)

    def test_blank_refresh_mode_override_clears_it(self):
        self.match.refresh_mode_override = RefreshMode.STOPPED
        self.match.save(update_fields=['refresh_mode_override'])
        response = self._post(refresh_mode_override='')
        self.assertRedirects(response, reverse('web:admin_matches'))
        self.match.refresh_from_db()
        self.assertIsNone(self.match.refresh_mode_override)


class AdminOddsAdjustmentViewTests(TestCase):
    """
    The unified user/user-match dashboard for the odds-adjustment and
    lay-spread cascades - previously only reachable via raw Django admin.
    Unlike Custom Odds' PricingOverride (one unified nullable-FK table),
    this data lives in three separate places (Match, User,
    UserMatchOddsOverride), so these tests specifically confirm the save
    dispatches to the right one AND that the change is visible through the
    real resolver functions, not just sitting in the right table.
    """
    def setUp(self):
        self.admin = User.objects.create_user(email='oddsadjadmin@example.com', password='testpass123', is_staff=True)
        self.client.force_login(self.admin)
        self.target_user = User.objects.create_user(email='oddsadjtarget@example.com', password='testpass123')
        self.sport = Sport.objects.create(name='OddsAdjSport', slug='odds-adj-sport')
        self.match = Match.objects.create(
            sport=self.sport, home_team='Home', away_team='Away',
            start_time=timezone.now() + timezone.timedelta(hours=1), status=Match.Status.SCHEDULED,
            odds_home=Decimal('2.00'), odds_draw=Decimal('3.00'), odds_away=Decimal('4.00'),
        )
        self.other_match = Match.objects.create(
            sport=self.sport, home_team='Home2', away_team='Away2',
            start_time=timezone.now() + timezone.timedelta(hours=1), status=Match.Status.SCHEDULED,
            odds_home=Decimal('5.00'), odds_draw=Decimal('6.00'), odds_away=Decimal('7.00'),
        )

    def _post(self, **overrides):
        data = {'action': 'save', 'user_id': '', 'match_id': '', 'adjustment': '', 'lay_spread_override': ''}
        data.update(overrides)
        return self.client.post(reverse('web:admin_odds_adjustment'), data)

    def test_match_only_save_writes_to_match(self):
        response = self._post(match_id=self.match.id, adjustment='0.30')
        self.assertRedirects(response, reverse('web:admin_odds_adjustment'))
        self.match.refresh_from_db()
        self.assertEqual(self.match.odds_adjustment, Decimal('0.30'))

    def test_user_only_save_writes_to_user(self):
        response = self._post(user_id=self.target_user.id, adjustment='0.75', lay_spread_override='0.20')
        self.assertRedirects(response, reverse('web:admin_odds_adjustment'))
        self.target_user.refresh_from_db()
        self.assertEqual(self.target_user.odds_adjustment_override, Decimal('0.75'))
        self.assertEqual(self.target_user.lay_spread_override, Decimal('0.20'))

    def test_user_and_match_save_creates_override_row(self):
        response = self._post(user_id=self.target_user.id, match_id=self.match.id, adjustment='1.10')
        self.assertRedirects(response, reverse('web:admin_odds_adjustment'))
        override = UserMatchOddsOverride.objects.get(user=self.target_user, match=self.match)
        self.assertEqual(override.adjustment, Decimal('1.10'))
        self.assertIsNone(override.lay_spread_override)

    def test_resaving_same_user_match_pair_updates_not_duplicates(self):
        UserMatchOddsOverride.objects.create(user=self.target_user, match=self.match, adjustment=Decimal('1.00'))
        self._post(user_id=self.target_user.id, match_id=self.match.id, adjustment='2.00')
        self.assertEqual(UserMatchOddsOverride.objects.filter(user=self.target_user, match=self.match).count(), 1)
        override = UserMatchOddsOverride.objects.get(user=self.target_user, match=self.match)
        self.assertEqual(override.adjustment, Decimal('2.00'))

    def test_no_user_and_no_match_rejected(self):
        response = self._post(adjustment='0.50')
        self.assertEqual(response.status_code, 200)
        self.match.refresh_from_db()
        self.assertIsNone(self.match.odds_adjustment)

    def test_no_values_rejected(self):
        response = self._post(user_id=self.target_user.id)
        self.assertEqual(response.status_code, 200)
        self.target_user.refresh_from_db()
        self.assertIsNone(self.target_user.odds_adjustment_override)

    def test_adjustment_out_of_range_rejected(self):
        response = self._post(user_id=self.target_user.id, adjustment='99.01')
        self.assertEqual(response.status_code, 200)
        self.target_user.refresh_from_db()
        self.assertIsNone(self.target_user.odds_adjustment_override)

    def test_lay_spread_out_of_range_rejected(self):
        response = self._post(user_id=self.target_user.id, lay_spread_override='6.00')
        self.assertEqual(response.status_code, 200)
        self.target_user.refresh_from_db()
        self.assertIsNone(self.target_user.lay_spread_override)

    def test_reset_user_clears_user_fields(self):
        self.target_user.odds_adjustment_override = Decimal('0.50')
        self.target_user.lay_spread_override = Decimal('0.30')
        self.target_user.save(update_fields=['odds_adjustment_override', 'lay_spread_override'])
        response = self.client.post(
            reverse('web:admin_odds_adjustment'), {'action': 'reset_user', 'user_id': self.target_user.id},
        )
        self.assertRedirects(response, reverse('web:admin_odds_adjustment'))
        self.target_user.refresh_from_db()
        self.assertIsNone(self.target_user.odds_adjustment_override)
        self.assertIsNone(self.target_user.lay_spread_override)

    def test_reset_match_clears_match_fields(self):
        self.match.odds_adjustment = Decimal('0.30')
        self.match.lay_spread_override = Decimal('0.20')
        self.match.save(update_fields=['odds_adjustment', 'lay_spread_override'])
        response = self.client.post(
            reverse('web:admin_odds_adjustment'), {'action': 'reset_match', 'match_id': self.match.id},
        )
        self.assertRedirects(response, reverse('web:admin_odds_adjustment'))
        self.match.refresh_from_db()
        self.assertIsNone(self.match.odds_adjustment)
        self.assertIsNone(self.match.lay_spread_override)

    def test_reset_user_match_deletes_row(self):
        override = UserMatchOddsOverride.objects.create(user=self.target_user, match=self.match, adjustment=Decimal('1.00'))
        response = self.client.post(
            reverse('web:admin_odds_adjustment'), {'action': 'reset_user_match', 'override_id': override.id},
        )
        self.assertRedirects(response, reverse('web:admin_odds_adjustment'))
        self.assertFalse(UserMatchOddsOverride.objects.filter(pk=override.id).exists())

    def test_end_to_end_user_only_visible_through_resolvers(self):
        self._post(user_id=self.target_user.id, adjustment='0.75', lay_spread_override='0.20')
        self.target_user.refresh_from_db()
        self.assertEqual(resolve_user_extra_adjustment(self.target_user, self.match), Decimal('0.75'))
        self.assertEqual(resolve_user_extra_lay_spread(self.target_user, self.match), Decimal('0.20'))
        # A different match for the same user still gets the user-global value.
        self.assertEqual(resolve_user_extra_adjustment(self.target_user, self.other_match), Decimal('0.75'))

    def test_end_to_end_user_match_beats_user_only(self):
        self._post(user_id=self.target_user.id, adjustment='0.10')
        self._post(user_id=self.target_user.id, match_id=self.match.id, adjustment='2.00')
        self.target_user.refresh_from_db()
        self.assertEqual(resolve_user_extra_adjustment(self.target_user, self.match), Decimal('2.00'))
        # The other match isn't covered by the specific override - falls back to user-global.
        self.assertEqual(resolve_user_extra_adjustment(self.target_user, self.other_match), Decimal('0.10'))

    def test_end_to_end_reset_falls_back_correctly(self):
        self._post(user_id=self.target_user.id, adjustment='0.10')
        self._post(user_id=self.target_user.id, match_id=self.match.id, adjustment='2.00')
        self.target_user.refresh_from_db()
        override = UserMatchOddsOverride.objects.get(user=self.target_user, match=self.match)
        self.client.post(
            reverse('web:admin_odds_adjustment'), {'action': 'reset_user_match', 'override_id': override.id},
        )
        self.assertEqual(resolve_user_extra_adjustment(self.target_user, self.match), Decimal('0.10'))

    def test_non_dashboard_user_redirected(self):
        self.client.logout()
        plain_user = User.objects.create_user(email='oddsadjplain@example.com', password='testpass123')
        self.client.force_login(plain_user)
        response = self.client.get(reverse('web:admin_odds_adjustment'))
        self.assertEqual(response.status_code, 302)

    def test_global_default_shows_no_warning_at_true_defaults(self):
        # Real bug found live: comparing a Decimal field to a bare float
        # literal in the template (`!= 0.10`) is never equal in Python even
        # at the true default, so the warning icon showed unconditionally.
        # Fixed by computing the comparison as a Decimal-to-Decimal check
        # in the view and passing a plain boolean to the template.
        response = self.client.get(reverse('web:admin_odds_adjustment'))
        self.assertTrue(response.context['global_is_default'])
        self.assertNotContains(response, 'Non-default global settings active')

    def test_global_default_shows_warning_when_changed(self):
        from apps.sports.models import OddsAdjustmentConfig
        config = OddsAdjustmentConfig.get_solo()
        config.default_adjustment = Decimal('0.50')
        config.save()
        response = self.client.get(reverse('web:admin_odds_adjustment'))
        self.assertFalse(response.context['global_is_default'])
        self.assertContains(response, 'Non-default global settings active')


class CasinoLaunchRedirectTests(TestCase):
    """
    Confirms casino_launch_view embeds the provider's returned URL in an
    iframe on our own page rather than redirecting the player's top-level
    tab there - per Waija's own docs ("open in an iframe or a new window")
    and confirmed with their support as what avoids a provider's own
    branded loading splash, with no server-side setting needed for it.
    """

    def setUp(self):
        CasinoConfig.objects.filter(pk=1).delete()
        self.config = CasinoConfig.objects.create(pk=1, is_enabled=True, max_launch_balance=Decimal('100.00'))
        self.brand = CasinoBrand.objects.create(brand_id=1, name='Pragmatic Play', is_active=True)
        self.game = CasinoGame.objects.create(
            game_id=737, game_uid='737', brand=self.brand, name='Aviator', category='slots', is_active=True,
        )
        self.user = User.objects.create_user(email='casinolaunchuser@example.com', password='testpass123')
        deposit_funds(self.user, Decimal('1000'))
        self.client.force_login(self.user)

    def test_launch_embeds_provider_url_in_iframe(self):
        # casino_launch_view (POST-only) redirects to the GET-safe
        # casino_play_view rather than rendering directly - Post/Redirect/
        # Get, so refreshing the play page never re-hits the POST-only
        # launch URL (which would 405 with a blank response otherwise).
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.post(reverse('web:casino_launch', args=[self.game.id]), {'amount': '10'}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<iframe')
        self.assertContains(response, 'https://mock-casino.local/play')

    def test_play_page_survives_a_refresh(self):
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            self.client.post(reverse('web:casino_launch', args=[self.game.id]), {'amount': '10'})
        # Simulate a page refresh: a plain GET on the same play URL the
        # browser ended up on, in a fresh request/response cycle.
        response = self.client.get(reverse('web:casino_play', args=[self.game.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<iframe')
        self.assertContains(response, 'https://mock-casino.local/play')

    def test_play_page_without_a_launch_redirects_to_lobby(self):
        response = self.client.get(reverse('web:casino_play', args=[self.game.id]))
        self.assertRedirects(response, reverse('web:casino_lobby'))


class CasinoBrowsingTests(TestCase):
    """Providers-first browsing: the lobby shows provider cards by default, a provider's own page never shows another provider's games, and search covers both."""

    def setUp(self):
        self.user = User.objects.create_user(email='casinobrowseuser@example.com', password='testpass123')
        self.client.force_login(self.user)
        self.brand_a = CasinoBrand.objects.create(brand_id=1, name='Pragmatic Play', is_active=True, game_count=1)
        self.brand_b = CasinoBrand.objects.create(brand_id=2, name='Evolution Live', is_active=True, game_count=1)
        self.game_a = CasinoGame.objects.create(
            game_id=1, game_uid='pragmaticslots/aviator-clone', brand=self.brand_a, name='Diamond Rush', is_active=True,
        )
        self.game_b = CasinoGame.objects.create(
            game_id=2, game_uid='evolution/roulette', brand=self.brand_b, name='Lightning Roulette', is_active=True,
        )

    def test_lobby_shows_provider_cards_not_games_by_default(self):
        response = self.client.get(reverse('web:casino_lobby'))
        self.assertContains(response, 'Pragmatic Play')
        self.assertContains(response, 'Evolution Live')
        self.assertNotContains(response, 'Diamond Rush')
        self.assertNotContains(response, 'Lightning Roulette')

    def test_provider_page_only_shows_that_providers_games(self):
        response = self.client.get(reverse('web:casino_provider', args=[self.brand_a.id]))
        self.assertContains(response, 'Diamond Rush')
        self.assertNotContains(response, 'Lightning Roulette')  # never mixed in from another provider

    def test_search_finds_both_matching_providers_and_games(self):
        response = self.client.get(reverse('web:casino_lobby'), {'q': 'Evolution'})
        self.assertContains(response, 'Evolution Live')  # provider name match
        response = self.client.get(reverse('web:casino_lobby'), {'q': 'Diamond'})
        self.assertContains(response, 'Diamond Rush')  # game name match
        self.assertNotContains(response, 'Lightning Roulette')


class SignupCountryCurrencyTests(TestCase):
    def test_pakistan_signup_gets_pkr_currency(self):
        response = self.client.post(reverse('web:register'), {
            'email': 'pkuser@example.com', 'password1': 'testpass123', 'password2': 'testpass123',
            'country': 'PK',
        })
        self.assertEqual(response.status_code, 302)
        user = get_user_model().objects.get(email='pkuser@example.com')
        self.assertEqual(user.country, 'PK')
        self.assertEqual(user.currency, 'PKR')

    def test_other_country_signup_gets_usd_currency(self):
        response = self.client.post(reverse('web:register'), {
            'email': 'ususer@example.com', 'password1': 'testpass123', 'password2': 'testpass123',
            'country': 'US',
        })
        self.assertEqual(response.status_code, 302)
        user = get_user_model().objects.get(email='ususer@example.com')
        self.assertEqual(user.country, 'US')
        self.assertEqual(user.currency, 'USD')

    def test_client_cannot_smuggle_a_currency_value(self):
        # currency is derived server-side from country only - even if a
        # tampered request tries to pass its own currency, it must be ignored.
        response = self.client.post(reverse('web:register'), {
            'email': 'sneaky@example.com', 'password1': 'testpass123', 'password2': 'testpass123',
            'country': 'US', 'currency': 'PKR',
        })
        self.assertEqual(response.status_code, 302)
        user = get_user_model().objects.get(email='sneaky@example.com')
        self.assertEqual(user.currency, 'USD')  # country US -> USD, the smuggled 'PKR' is ignored

    def test_invalid_country_rejected(self):
        response = self.client.post(reverse('web:register'), {
            'email': 'badcountry@example.com', 'password1': 'testpass123', 'password2': 'testpass123',
            'country': 'ZZ',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'valid country')
        self.assertFalse(get_user_model().objects.filter(email='badcountry@example.com').exists())


class CasinoCurrencyBannerTests(TestCase):
    """The PKR rate/balance/disclaimer banner shown across every casino page - see templates/web/_casino_currency_banner.html."""

    def setUp(self):
        from apps.wallet.models import FxRateConfig

        FxRateConfig.objects.update_or_create(pk=1, defaults={'usd_pkr_rate': Decimal('280.0000')})
        self.pkr_user = User.objects.create_user(email='casinobannerpkr@example.com', password='testpass123', currency='PKR')
        deposit_funds(self.pkr_user, Decimal('53'))
        self.usd_user = User.objects.create_user(email='casinobannerusd@example.com', password='testpass123', currency='USD')
        deposit_funds(self.usd_user, Decimal('53'))
        self.brand = CasinoBrand.objects.create(brand_id=1, name='Pragmatic Play', is_active=True, game_count=1)

    def test_pkr_user_sees_rate_and_dual_currency_balance(self):
        self.client.force_login(self.pkr_user)
        response = self.client.get(reverse('web:casino_lobby'))
        self.assertContains(response, '1 USD = Rs 280.0000')
        self.assertContains(response, 'Rs 53.00')
        self.assertContains(response, '$0.19')  # 53 / 280 = 0.18928571..., quantized to 0.19
        self.assertContains(response, 'settled in US Dollars')

    def test_usd_user_sees_no_banner(self):
        self.client.force_login(self.usd_user)
        response = self.client.get(reverse('web:casino_lobby'))
        self.assertNotContains(response, 'USD Rate:')
        self.assertNotContains(response, 'settled in US Dollars')

    def test_banner_also_shows_on_provider_page(self):
        self.client.force_login(self.pkr_user)
        response = self.client.get(reverse('web:casino_provider', args=[self.brand.id]))
        self.assertContains(response, 'settled in US Dollars')


class CasinoGameBlurTests(TestCase):
    """
    Admin-drawn per-game thumbnail blur region (a face/clothing area in
    suggestive cover art, say) - see admin_casino_game_blur_view and
    templates/web/_casino_game_card.html. Stored as percentages of the
    image so it's independent of wherever/however large the thumbnail
    actually renders.
    """

    def setUp(self):
        self.admin = User.objects.create_user(email='bluradmin@example.com', password='testpass123', is_staff=True)
        self.player = User.objects.create_user(email='blurplayer@example.com', password='testpass123')
        self.brand = CasinoBrand.objects.create(brand_id=1, name='Bgaming', is_active=True)
        self.game = CasinoGame.objects.create(
            game_id=1, game_uid='bgaming/some-game', brand=self.brand, name='Some Game', is_active=True,
            logo_url='https://example.com/thumb.jpg',
        )

    def test_admin_can_save_a_blur_region(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('web:admin_casino_game_blur', args=[self.game.id]),
            {'top': '10.5', 'left': '20', 'width': '30', 'height': '25.25'},
        )
        self.assertRedirects(response, reverse('web:admin_casino_games'))
        self.game.refresh_from_db()
        self.assertEqual(self.game.blur_region, {'top': 10.5, 'left': 20.0, 'width': 30.0, 'height': 25.25})

    def test_admin_can_clear_a_blur_region(self):
        self.game.blur_region = {'top': 1, 'left': 2, 'width': 3, 'height': 4}
        self.game.save(update_fields=['blur_region'])
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('web:admin_casino_game_blur', args=[self.game.id]), {'action': 'clear'},
        )
        self.assertRedirects(response, reverse('web:admin_casino_games'))
        self.game.refresh_from_db()
        self.assertIsNone(self.game.blur_region)

    def test_a_zero_size_box_is_rejected(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('web:admin_casino_game_blur', args=[self.game.id]),
            {'top': '10', 'left': '10', 'width': '0', 'height': '0'},
        )
        self.assertEqual(response.status_code, 200)
        self.game.refresh_from_db()
        self.assertIsNone(self.game.blur_region)

    def test_non_admin_cannot_reach_the_blur_editor(self):
        self.client.force_login(self.player)
        response = self.client.get(reverse('web:admin_casino_game_blur', args=[self.game.id]))
        self.assertNotEqual(response.status_code, 200)

    def test_game_card_renders_blur_overlay_only_when_region_is_set(self):
        self.client.force_login(self.player)
        response = self.client.get(reverse('web:casino_provider', args=[self.brand.id]))
        self.assertNotContains(response, 'backdrop-filter')

        self.game.blur_region = {'top': 10, 'left': 20, 'width': 30, 'height': 40}
        self.game.save(update_fields=['blur_region'])
        response = self.client.get(reverse('web:casino_provider', args=[self.brand.id]))
        self.assertContains(response, 'backdrop-filter')
        self.assertContains(response, 'top:10%')
        self.assertContains(response, 'left:20%')


class CasinoGameLiveBlurTests(TestCase):
    """
    Admin-positioned FIXED overlay box shown for the whole session on
    casino_play.html (a live dealer's face/upper costume, say) - see
    admin_casino_game_live_blur_view. Deliberately NOT real-time tracking
    (the video is inside a cross-origin iframe our page cannot read
    pixels from), so this is a manually-drawn fixed box, positioned
    against the free demo stream so the admin can see real footage.
    """

    def setUp(self):
        CasinoConfig.objects.filter(pk=1).delete()
        CasinoConfig.objects.create(pk=1, is_enabled=True, max_launch_balance=Decimal('100.00'))
        self.admin = User.objects.create_user(email='liveblur_admin@example.com', password='testpass123', is_staff=True)
        self.player = User.objects.create_user(email='liveblur_player@example.com', password='testpass123')
        deposit_funds(self.player, Decimal('1000'))
        self.brand = CasinoBrand.objects.create(brand_id=1, name='Evolution Live', is_active=True)
        self.game = CasinoGame.objects.create(
            game_id=1, game_uid='evolution/some-live-game', brand=self.brand, name='Some Live Game', is_active=True,
            logo_url='https://example.com/thumb.jpg',
        )

    def test_editor_loads_the_real_demo_stream(self):
        self.client.force_login(self.admin)
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.get(reverse('web:admin_casino_game_live_blur', args=[self.game.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'https://mock-casino.local/demo')

    def test_admin_can_save_a_live_blur_box(self):
        self.client.force_login(self.admin)
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.post(
                reverse('web:admin_casino_game_live_blur', args=[self.game.id]),
                {'top': '5', 'left': '30', 'width': '40', 'height': '35'},
            )
        self.assertRedirects(response, reverse('web:admin_casino_games'))
        self.game.refresh_from_db()
        self.assertEqual(self.game.live_blur_region, {'top': 5.0, 'left': 30.0, 'width': 40.0, 'height': 35.0})

    def test_admin_can_clear_a_live_blur_box(self):
        self.game.live_blur_region = {'top': 1, 'left': 2, 'width': 3, 'height': 4}
        self.game.save(update_fields=['live_blur_region'])
        self.client.force_login(self.admin)
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.post(
                reverse('web:admin_casino_game_live_blur', args=[self.game.id]), {'action': 'clear'},
            )
        self.assertRedirects(response, reverse('web:admin_casino_games'))
        self.game.refresh_from_db()
        self.assertIsNone(self.game.live_blur_region)

    def test_non_admin_cannot_reach_the_live_blur_editor(self):
        self.client.force_login(self.player)
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.get(reverse('web:admin_casino_game_live_blur', args=[self.game.id]))
        self.assertNotEqual(response.status_code, 200)

    def test_play_page_shows_persistent_overlay_only_when_region_is_set(self):
        self.client.force_login(self.player)
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.post(reverse('web:casino_launch', args=[self.game.id]), {'amount': '10'}, follow=True)
        self.assertNotContains(response, 'backdrop-filter:blur(18px)')

        self.game.live_blur_region = {'top': 5, 'left': 30, 'width': 40, 'height': 35}
        self.game.save(update_fields=['live_blur_region'])
        with patch('apps.casino.services.get_casino_provider', return_value=MockCasinoProvider()):
            response = self.client.post(reverse('web:casino_launch', args=[self.game.id]), {'amount': '10'}, follow=True)
        self.assertContains(response, 'backdrop-filter:blur(18px)')
        self.assertContains(response, 'top:5%')
        self.assertContains(response, 'left:30%')


class CasinoGameLiveBlurDemoFallbackTests(TestCase):
    """Most live-dealer games have no demo mode at all - the editor must fall back to the thumbnail, not crash or dead-end."""

    def setUp(self):
        self.admin = User.objects.create_user(email='liveblur_fallback_admin@example.com', password='testpass123', is_staff=True)
        self.brand = CasinoBrand.objects.create(brand_id=1, name='Evolution Live', is_active=True)
        self.game = CasinoGame.objects.create(
            game_id=1, game_uid='evolution/some-live-game', brand=self.brand, name='Some Live Game', is_active=True,
            logo_url='https://example.com/thumb.jpg',
        )
        self.client.force_login(self.admin)

    def _no_demo_provider(self):
        provider = MockCasinoProvider()
        provider.fetch_demo_url = lambda game_uid, return_url='': (_ for _ in ()).throw(
            RuntimeError('Demo mode is not available for this game.')
        )
        return provider

    def test_falls_back_to_thumbnail_when_demo_unavailable(self):
        with patch('apps.casino.services.get_casino_provider', return_value=self._no_demo_provider()):
            response = self.client.get(reverse('web:admin_casino_game_live_blur', args=[self.game.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'https://example.com/thumb.jpg')
        self.assertContains(response, 'Demo mode is not available')

    def test_can_still_save_a_box_using_the_thumbnail_fallback(self):
        with patch('apps.casino.services.get_casino_provider', return_value=self._no_demo_provider()):
            response = self.client.post(
                reverse('web:admin_casino_game_live_blur', args=[self.game.id]),
                {'top': '5', 'left': '30', 'width': '40', 'height': '35'},
            )
        self.assertRedirects(response, reverse('web:admin_casino_games'))
        self.game.refresh_from_db()
        self.assertEqual(self.game.live_blur_region, {'top': 5.0, 'left': 30.0, 'width': 40.0, 'height': 35.0})
